#!/usr/bin/env python3
"""PPTX スライド規約の機械チェック。

使い方:
  python3 scripts/check_deck.py out.pptx
  python3 scripts/check_deck.py slide-spec/super_template.pptx --template

正典: references/slide-rules.md
終了コード: FAIL があれば 1。python-pptx 必須。
"""
import re
import sys
import zipfile
from pathlib import Path

# 自ブランドで禁止するレガシー色があればここに列挙（HEX 6桁・#なし）
OLD_COLORS = []
TITLE_MAX = 40
EMU_W, EMU_H = 12192000, 6858000

fails, warns = [], []


def fail(msg):
    fails.append(msg)


def warn(msg):
    warns.append(msg)


# テンプレ集そのものを検査するときだけ True（--template）。
# 通常のデッキでは「Text N」「◯◯」等のプレースホルダーが残っていたら FAIL にする。
TEMPLATE_MODE = False

# 表記ゆれ代表ペア（slide-rules §7.6: 1資料1用語）。両方の表記が同一資料に現れたら WARN。
TERM_VARIANTS = [
    ("メモリー", r"メモリー", "メモリ", r"メモリ(?!ー)(?!・CPU)"),
    ("紐付", r"紐付", "紐づ", r"紐づ"),
    ("ユーザー", r"ユーザー", "利用者", r"利用者"),
    ("アセット", r"アセット", "資産", r"資産"),
    ("フォルダー", r"フォルダー", "フォルダ", r"フォルダ(?!ー)"),
    ("サーバー", r"サーバー", "サーバ", r"サーバ(?!ー)"),
    ("メンバー", r"メンバー", "メンバ", r"メンバ(?!ー)"),
    ("コンピューター", r"コンピューター", "コンピュータ", r"コンピュータ(?!ー)"),
    ("問い合わせ", r"問い合わせ", "問合せ", r"問合せ"),
]


# AI臭ワード（slide-rules §7.9 / references/ai-smell-lexicon.md）。高確度語のみ WARN。
AI_SMELL_WORDS = [
    "まさに", "非常に", "極めて", "圧倒的", "画期的", "革新的", "次世代の",
    "過言ではありません", "に他なりません", "シームレス", "シナジー", "ソリューション",
    "エンドツーエンド", "ブラッシュアップ", "付加価値",
    "寄り添い", "伴走し", "二人三脚", "さらなる高みへ", "邁進",
    "昨今", "変化の激しい", "という点において", "の観点から",
    "させていただきます", "いただけますと幸いです",
    "と言えるでしょう", "と考えられます", "することが可能です",
]


def check_ai_smell(pages):
    """slide-rules §7.9: AI臭の高確度語を検出（WARN）"""
    hits = {}
    for i, txt in pages:
        found = [w for w in AI_SMELL_WORDS if w in txt]
        if found:
            hits[i] = found
    if hits:
        detail = "、".join(f"p{i}:「{'/'.join(ws[:3])}」" for i, ws in sorted(hits.items())[:6])
        warn(f"AI臭ワード検出: {detail}（§7.9 / ai-smell-lexicon.md。素の動詞・直球の言い方に置き換え）")
    dash_pages = sorted({i for i, txt in pages if " — " in txt or "—" in txt})
    if dash_pages:
        warn(f"ダッシュ「 — 」連結 p{dash_pages}（AI文体の典型。句点・「：」・括弧に置き換え。§7.9）")


def check_terms(pages):
    """pages: [(idx, text), ...] 資料全体で両方の表記が出たら WARN"""
    for la, pa, lb, pb in TERM_VARIANTS:
        hits_a = sorted({i for i, t in pages if re.search(pa, t)})
        hits_b = sorted({i for i, t in pages if re.search(pb, t)})
        if hits_a and hits_b:
            warn(f"表記ゆれ疑い: 「{la}」p{hits_a} と「{lb}」p{hits_b} が混在（§7.6 1資料1用語。別概念なら可・目視確認）")


def check_title(idx, title, explicit_break=False):
    t = title.strip()
    if not t:
        warn(f"p{idx}: タイトルが空（表紙/扉なら可）")
        return
    if re.search(r"(です|ます|でした|ました)[。．.]?$", t):
        fail(f"p{idx}: タイトルがですます調 → 体言止めに: 「{t}」")
    # 2行まで許容。1行目安(TITLE_MAX=40字)超は WARN、>80字は FAIL。文字縮小で1行詰めは不可（§2.1）
    tlen = sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)
    if tlen > TITLE_MAX * 2:
        fail(f"p{idx}: タイトル 全角換算{tlen:.0f}字（>{TITLE_MAX*2}・2行にも収まらない。主張を絞る）: 「{t}」")
    elif tlen > TITLE_MAX and not explicit_break:
        warn(f"p{idx}: タイトル 全角換算{tlen:.0f}字（2行になる想定。意味の切れ目で改行・泣き別れなし・文字縮小で1行に詰めない）: 「{t}」")
    if re.match(r"^(Step|STEP|ステップ)\s*\d", t):
        fail(f"p{idx}: タイトルに Step 連結（タグチップで表現）: 「{t}」")
    if re.search(r"^(この|その|ここまで)", t):
        fail(f"p{idx}: 他スライド参照語で始まるタイトル: 「{t}」")
    if t.count("（") + t.count("(") >= 2:
        warn(f"p{idx}: タイトルに丸括弧が多い: 「{t}」")
    if not TEMPLATE_MODE and re.search(r"[◯○]{2,}|Text\s*\d|ラベル\s*\d|タイトル\s*\d|Source\s*\d|YYYY|ダミー|^資料名$|^会社名$", t):
        fail(f"p{idx}: テンプレのプレースホルダーが残っている（タイトルはストーリーラインから書く — slide-rules §2.8）: 「{t}」")


def check_title_variety(titles):
    real = [t.strip() for t in titles if t and t.strip()]
    if len(real) < 5:
        return
    molds = {
        "「◯◯は、…」": lambda t: re.match(r"^.{1,12}は[、,]", t),
        "「◯◯には、…がある」": lambda t: re.search(r"には[、,].*(ある|存在する)$", t),
        "数を主語にした形（「3つの…」）": lambda t: re.match(r"^[0-9０-９一二三四五六七八九十]+つ", t),
    }
    for name, f in molds.items():
        hits = [t for t in real if f(t)]
        if len(hits) >= max(4, int(len(real) * 0.6)):
            warn(f"タイトルの文型が {name} に偏っている（{len(hits)}/{len(real)}枚）。テンプレの見本文型をなぞらず、主張ごとに自然な文で書く — slide-rules §2.8")


KANSUJI = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
           "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
COUNT_WORDS = r"(?:段階|フェーズ|ステップ|柱|論点)"
ENUM_FAMILY = [
    (r"段階|フェーズ|ステップ", r"(?:フェーズ|ステップ|STEP|Step|PHASE|Phase|段階)"),
    (r"柱", r"(?:柱)"),
    (r"論点", r"(?:論点)"),
]


def _to_int(s):
    s = s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))
    if s.isdigit():
        return int(s)
    return KANSUJI.get(s)


def check_count_match(idx, title, text):
    """タイトルの「N段階」等が本文の連番ラベルと食い違っていたら FAIL"""
    if not title or not text:
        return
    for m in re.finditer(r"([0-9０-９]+|[一二三四五六七八九十])\s*(" + COUNT_WORDS + ")", title):
        n = _to_int(m.group(1))
        if not n or not (2 <= n <= 12):
            continue
        fam = next((lab for pat, lab in ENUM_FAMILY if re.search(pat, m.group(2))), None)
        if not fam:
            continue
        found = [v for v in (_to_int(x.group(1))
                             for x in re.finditer(fam + r"\s*([0-9０-９]+)", text)) if v]
        if found and max(found) != n:
            fail(f"p{idx}: タイトルの「{m.group(0)}」と本文の連番（最大{max(found)}）が食い違う"
                 f"（§2.9 数はページの中身と一致させる）: 「{title}」")


def _fwlen(t):
    return sum(0.5 if ord(ch) < 0x3000 else 1 for ch in t)


def check_orphan(idx, shape, label="タイトル"):
    """箱幅とフォントサイズから1行容量を推定し、最終行が1〜3字だけになる折返し（泣き別れ）を WARN"""
    try:
        if not (shape.has_text_frame and shape.width):
            return
        tf = shape.text_frame
        raw = tf.text
        if not raw.strip():
            return
        szs = [r.font.size.pt for p in tf.paragraphs for r in p.runs if r.font.size]
        if not szs:
            return
        pt = max(szs)
        w_in = shape.width / 914400.0
        cap = max(4, int(w_in / (pt / 72.0)))
        for line in raw.split("\n"):
            L = _fwlen(line.strip())
            if L <= cap:
                continue
            last = L - cap * ((int(-(-L // cap))) - 1)
            if 0 < last <= 3:
                warn(f"p{idx}: {label}が泣き別れの疑い（推定{int(-(-L // cap))}行目が約{last:.0f}字だけ。意味の切れ目で明示改行 — slide-rules §2.13）: 「{line.strip()[:40]}」")
    except Exception:
        return


def check_fill_ratio(idx, slide, title_shape):
    """本文（タイトル下〜出典行上）の縦幅に対する本文要素の占有率。55%未満は WARN（§5.13）"""
    try:
        top_lim = 1.6 * 914400
        bot_lim = 6.8 * 914400
        tops, bots = [], []
        for sh in slide.shapes:
            if sh is title_shape or sh.top is None or sh.height is None:
                continue
            t, b = sh.top, sh.top + sh.height
            if b <= top_lim or t >= bot_lim:
                continue
            if sh.height < 914400 * 0.02 and sh.width and sh.width > 914400 * 11:
                continue
            tops.append(max(t, top_lim))
            bots.append(min(b, bot_lim))
        if not tops:
            return
        cov = (max(bots) - min(tops)) / (bot_lim - top_lim)
        if cov < 0.55:
            warn(f"p{idx}: 版面充填率 {cov*100:.0f}%（本文が版面の半分未満。情報を足すか型を変える。飾りで埋めない — slide-rules §5.13）")
    except Exception:
        return


def check_pptx(path):
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        sys.exit("python-pptx が必要: pip3 install python-pptx")
    prs = Presentation(path)
    if abs(prs.slide_width - EMU_W) > 2000 or abs(prs.slide_height - EMU_H) > 2000:
        fail(f"スライドサイズ {prs.slide_width}x{prs.slide_height} ≠ 16:9 {EMU_W}x{EMU_H}")

    titles = []
    term_pages = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
        for n in sorted(names, key=lambda s: int(re.search(r"slide(\d+)", s).group(1))):
            xml = z.read(n).decode("utf8", "ignore")
            idx = int(re.search(r"slide(\d+)", n).group(1))
            term_pages.append((idx, " ".join(re.findall(r"<a:t>(.*?)</a:t>", xml, re.S))))
            rr = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                if 'prst="roundRect"' in sp:
                    h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', sp)
                    if h and int(h.group(1)) >= 365760:
                        rr += 1
            if rr:
                fail(f"p{idx}: roundRect の大きなボックス ×{rr}（直角 rect に）")
            fb = 0
            for m in re.finditer(r"<p:sp>.*?</p:sp>", xml, re.S):
                sp = m.group(0)
                spPr = re.search(r"<p:spPr>.*?</p:spPr>", sp, re.S)
                if not spPr:
                    continue
                pr = spPr.group(0)
                h = re.search(r'<a:ext cx="\d+" cy="(\d+)"', pr)
                if not (h and int(h.group(1)) >= 365760):
                    continue
                ln = re.search(r"<a:ln[ >].*?</a:ln>", pr, re.S)
                filled = "<a:solidFill>" in re.sub(r"<a:ln[ >].*?</a:ln>", "", pr, flags=re.S)
                if filled and ln and "<a:solidFill>" in ln.group(0):
                    fb += 1
            if fb:
                warn(f"p{idx}: 塗りあり図形に枠線 ×{fb}（塗りカードは line なし — slide-rules §5.3）")
            for c in OLD_COLORS:
                if f'val="{c}"' in xml or f'val="{c.lower()}"' in xml:
                    fail(f"p{idx}: 禁止色 {c}")
            lb = len(re.findall(r"<a:t>\s*•", xml))
            if lb:
                fail(f"p{idx}: ブレット「•」を本文テキストに直打ち ×{lb}（buChar/buNone 書式に — slide-rules §7.3）")
            ld = len(re.findall(r"<a:t>\s*[–‐-]\s\s", xml))
            if ld:
                warn(f"p{idx}: 第2階層マーカー「– 」らしきテキスト直打ち ×{ld}（buChar 書式に — slide-rules §7.3）")
        for n in z.namelist():
            if "theme" in n or "slideMaster" in n or "slideLayout" in n:
                xml = z.read(n).decode("utf8", "ignore")
                for c in OLD_COLORS:
                    if f'val="{c}"' in xml:
                        warn(f"{n}: 禁止色 {c}")

    for i, s in enumerate(prs.slides, 1):
        title = ""
        if s.shapes.title is not None and s.shapes.title.has_text_frame:
            title = s.shapes.title.text_frame.text
        else:
            cands = []
            for sh in s.shapes:
                if sh.has_text_frame and sh.top is not None and sh.top < Emu(1097280):
                    sz = max((r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size), default=0)
                    cands.append((sz, sh.text_frame.text))
            if cands:
                title = max(cands)[1]
        title_shape = None
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text == title:
                title_shape = sh
                break
        if title_shape is not None:
            check_orphan(i, title_shape, "タイトル")
        if i == 1:
            for sh in s.shapes:
                if sh is not title_shape and sh.has_text_frame:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) >= 13 and len(sh.text_frame.text) < 80:
                        check_orphan(i, sh, "表紙の文字")
        else:
            check_fill_ratio(i, s, title_shape)
        explicit_break = "\n" in title.strip() and all(_fwlen(l.strip()) <= TITLE_MAX for l in title.split("\n"))
        title = title.replace("\n", " ")
        titles.append(title)
        check_title(i, title, explicit_break)
        check_count_match(i, title, dict(term_pages).get(i, ""))
        for sh in s.shapes:
            if sh.has_text_frame and sh.top is not None and Emu(1097280) <= sh.top < Emu(1600200):
                txt = sh.text_frame.text.strip()
                if txt and "\n" not in txt and len(txt) < 60 and txt != title:
                    szs = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
                    if szs and max(szs) <= 12 and not any(r.font.bold for p in sh.text_frame.paragraphs for r in p.runs):
                        warn(f"p{i}: タイトル直下にサブタイトルらしき行: 「{txt}」")
    check_terms(term_pages)
    check_ai_smell(term_pages)
    return titles


def main():
    global TEMPLATE_MODE
    args = sys.argv[1:]
    if "--template" in args:
        TEMPLATE_MODE = True
        args.remove("--template")
    if not args:
        sys.exit(__doc__)
    p = args[0]
    if not p.lower().endswith(".pptx"):
        sys.exit("PPTX のみ対応: python3 scripts/check_deck.py deck.pptx")
    if not Path(p).is_file():
        sys.exit(f"ファイルが見つからない: {p}")
    titles = check_pptx(p)
    if not TEMPLATE_MODE:
        check_title_variety(titles)
    print("=== タイトル一覧（上から通し読みしてストーリーが繋がるか確認） ===")
    for i, t in enumerate(titles, 1):
        print(f"{i:>3}  {t or '(なし)'}")
    print()
    for w in warns:
        print("WARN ", w)
    for f in fails:
        print("FAIL ", f)
    print(f"\n{len(fails)} FAIL / {len(warns)} WARN")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
